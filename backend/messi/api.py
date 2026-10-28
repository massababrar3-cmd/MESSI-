from __future__ import annotations
import io, os, csv, json, time, secrets, copy
from pathlib import Path
from collections import defaultdict, deque
from contextlib import asynccontextmanager
from zoneinfo import ZoneInfo
from fastapi import FastAPI, Request, HTTPException, UploadFile, File
from fastapi.responses import JSONResponse, StreamingResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field, ValidationError, ConfigDict
from sqlalchemy import select, update
from sqlalchemy.exc import IntegrityError
from .db import (
    Session,
    initialize,
    Project,
    Member,
    User,
    LoginSession,
    Scenario,
    Run,
    Audit,
    audit,
    now,
    ROOT,
)
from .domain import Inputs, DomainError, parse_csv, parse_obj, input_hash, sample_inputs
from .scheduler import graph_issues, validate_result
from .security import identity, hash_password, verify_password, digest

ALLOWED = set(
    os.environ.get(
        "MESSI_ORIGINS",
        "http://127.0.0.1:5173,http://localhost:5173,http://127.0.0.1:8000,http://localhost:8000",
    ).split(",")
)
SECURE = os.environ.get("MESSI_SECURE_COOKIES", "0") == "1"
login_attempts = defaultdict(deque)
DUMMY = hash_password("non-account-timing-placeholder")


@asynccontextmanager
async def lifespan(app):
    initialize()
    yield


app = FastAPI(title="MESSI Mine Planning API", version="0.1.0", lifespan=lifespan)


@app.middleware("http")
async def guards(request, call_next):
    if request.method in ("POST", "PUT", "PATCH", "DELETE"):
        if request.headers.get("origin") not in ALLOWED:
            return JSONResponse({"detail": "Origin not allowed"}, status_code=403)
        try:
            size = int(request.headers.get("content-length", "0"))
        except ValueError:
            return JSONResponse({"detail": "Invalid Content-Length"}, status_code=400)
        if size > 5_000_000:
            return JSONResponse({"detail": "Request exceeds 5 MB limit"}, status_code=413)
        # Enforce the real streamed byte count, including chunked requests.
        body = bytearray()
        async for chunk in request.stream():
            if len(body) + len(chunk) > 5_000_000:
                return JSONResponse({"detail": "Request exceeds 5 MB limit"}, status_code=413)
            body.extend(chunk)
        request._body = bytes(body)
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "same-origin"
    response.headers["Content-Security-Policy"] = (
        "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data: blob:; connect-src 'self'; worker-src 'self' blob:; object-src 'none'; base-uri 'self'; frame-ancestors 'none'"
    )
    if request.url.path.startswith("/api/"):
        response.headers["Cache-Control"] = "no-store"
    return response


@app.exception_handler(DomainError)
async def domain_error(request, exc):
    return JSONResponse({"detail": str(exc)}, status_code=422)


@app.exception_handler(ValidationError)
async def validation_error(request, exc):
    return JSONResponse(
        {
            "detail": "; ".join(
                f"{'.'.join(str(x) for x in e['loc'])}: {e['msg']}" for e in exc.errors()
            )
        },
        status_code=422,
    )


def access(db, req, project_id, edit=False, approve=False, owner=False):
    user, _ = identity(db, req, req.method not in ["GET", "HEAD"])
    member = db.scalar(
        select(Member).where(Member.project_id == project_id, Member.user_id == user.id)
    )
    if not member:
        raise HTTPException(404, "Project not found")
    if (
        (owner and member.role != "owner")
        or (edit and member.role not in ["owner", "planner"])
        or (approve and member.role not in ["owner", "reviewer"])
    ):
        raise HTTPException(403, "Your project role does not allow this action")
    return user, member


def get_scenario(db, req, sid, edit=False):
    s = db.get(Scenario, sid)
    if not s:
        raise HTTPException(404, "Scenario not found")
    user, member = access(db, req, s.project_id, edit=edit)
    return s, user, member


def scenario_data(s):
    return {
        "id": s.id,
        "project_id": s.project_id,
        "name": s.name,
        "revision": s.revision,
        "parent_id": s.parent_id,
        "inputs": s.inputs,
    }


def run_data(r, full=True):
    d = {
        k: getattr(r, k)
        for k in [
            "id",
            "project_id",
            "scenario_id",
            "revision",
            "created",
            "input_hash",
            "status",
            "progress",
            "error",
            "approved_by",
            "approved_at",
        ]
    }
    if full:
        d.update(result=r.result, snapshot=r.snapshot)
    return d


def replace_inputs(db, s, user, revision, inputs, action):
    canonical = Inputs.model_validate(inputs).model_dump()
    changed = db.execute(
        update(Scenario)
        .where(Scenario.id == s.id, Scenario.revision == revision)
        .values(inputs=canonical, revision=revision + 1)
    )
    if changed.rowcount != 1:
        raise HTTPException(409, "Scenario changed. Reload before applying your edit.")
    audit(db, s.project_id, user.name, action, f"{s.id} revision {revision + 1}")
    db.flush()
    db.expire(s)
    return scenario_data(s)


class Login(BaseModel):
    email: str = Field(max_length=200)
    password: str = Field(max_length=256)


@app.get("/api/health")
def health():
    return {
        "status": "ok",
        "version": "0.1.0",
        "engine": "deterministic baseline",
        "deployment": "development pilot",
    }


@app.post("/api/auth/login")
def login(body: Login, request: Request):
    key = request.client.host if request.client else "unknown"
    attempts = login_attempts[key]
    stamp = time.time()
    while attempts and attempts[0] < stamp - 300:
        attempts.popleft()
    if len(attempts) >= 20:
        raise HTTPException(429, "Too many attempts; wait five minutes")
    attempts.append(stamp)
    with Session.begin() as db:
        user = db.scalar(select(User).where(User.email == body.email.strip().lower()))
        if not verify_password(body.password, user.password if user else DUMMY) or not user:
            raise HTTPException(401, "Incorrect email or password")
        token = secrets.token_urlsafe(40)
        csrf = secrets.token_urlsafe(32)
        db.add(
            LoginSession(
                token_hash=digest(token), user_id=user.id, csrf=csrf, expires=time.time() + 8 * 3600
            )
        )
        response = JSONResponse(
            {"user": {"id": user.id, "name": user.name, "email": user.email}, "csrf": csrf}
        )
        response.set_cookie(
            "messi_session",
            token,
            httponly=True,
            secure=SECURE,
            samesite="lax",
            max_age=8 * 3600,
            path="/",
        )
        return response


@app.get("/api/auth/me")
def me(request: Request):
    with Session() as db:
        u, s = identity(db, request)
        return {"user": {"id": u.id, "name": u.name, "email": u.email}, "csrf": s.csrf}


@app.post("/api/auth/logout")
def logout(request: Request):
    with Session.begin() as db:
        _, s = identity(db, request, True)
        db.delete(s)
    response = JSONResponse({"ok": True})
    response.delete_cookie("messi_session")
    return response


class NewProject(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    crs: str = Field(min_length=1, max_length=160)
    datum: str = Field(min_length=1, max_length=160)
    timezone: str = "Australia/Brisbane"
    sample: bool = False


@app.get("/api/projects")
def projects(request: Request):
    with Session() as db:
        u, _ = identity(db, request)
        pairs = db.execute(
            select(Project, Member.role)
            .join(Member, Member.project_id == Project.id)
            .where(Member.user_id == u.id)
            .order_by(Project.created)
        ).all()
        return [
            {
                "id": p.id,
                "name": p.name,
                "crs": p.crs,
                "datum": p.datum,
                "timezone": p.timezone,
                "role": role,
                "created": p.created,
            }
            for p, role in pairs
        ]


@app.post("/api/projects", status_code=201)
def new_project(body: NewProject, request: Request):
    try:
        ZoneInfo(body.timezone)
    except Exception:
        raise HTTPException(422, "Unknown IANA timezone")
    with Session.begin() as db:
        u, _ = identity(db, request, True)
        p = Project(name=body.name, crs=body.crs, datum=body.datum, timezone=body.timezone)
        db.add(p)
        db.flush()
        db.add(Member(project_id=p.id, user_id=u.id, role="owner"))
        inputs = sample_inputs()
        inputs["timezone"] = p.timezone
        if not body.sample:
            inputs["blocks"] = []
            inputs["source"] = {}
        s = Scenario(project_id=p.id, name="Baseline plan", inputs=inputs)
        db.add(s)
        db.flush()
        audit(db, p.id, u.name, "project.created", p.name)
        return {"id": p.id, "scenario_id": s.id}


@app.get("/api/projects/{pid}/scenarios")
def scenarios(pid: str, request: Request):
    with Session() as db:
        access(db, request, pid)
        return [
            scenario_data(s)
            for s in db.scalars(
                select(Scenario).where(Scenario.project_id == pid).order_by(Scenario.name)
            )
        ]


@app.get("/api/scenarios/{sid}")
def scenario(sid: str, request: Request):
    with Session() as db:
        return scenario_data(get_scenario(db, request, sid)[0])


class EditScenario(BaseModel):
    revision: int = Field(ge=1)
    inputs: Inputs


@app.put("/api/scenarios/{sid}")
def edit_scenario(sid: str, body: EditScenario, request: Request):
    with Session.begin() as db:
        s, u, _ = get_scenario(db, request, sid, True)
        return replace_inputs(db, s, u, body.revision, body.inputs.model_dump(), "scenario.edited")


class Branch(BaseModel):
    name: str = Field(min_length=1, max_length=100)
    revision: int = Field(ge=1)


@app.post("/api/scenarios/{sid}/branch", status_code=201)
def branch(sid: str, body: Branch, request: Request):
    with Session.begin() as db:
        s, u, _ = get_scenario(db, request, sid, True)
        if s.revision != body.revision:
            raise HTTPException(409, "Scenario changed; reload before branching")
        new = Scenario(
            project_id=s.project_id, name=body.name, parent_id=s.id, inputs=copy.deepcopy(s.inputs)
        )
        db.add(new)
        db.flush()
        audit(db, s.project_id, u.name, "scenario.branched", f"{s.id} → {new.id}")
        return scenario_data(new)


@app.post("/api/scenarios/{sid}/import-preview")
async def preview(sid: str, request: Request, file: UploadFile = File(...), units: str = "m"):
    with Session() as db:
        get_scenario(db, request, sid, True)
    content = await file.read(2_000_001)
    if len(content) > 2_000_000:
        raise HTTPException(413, "CSV exceeds 2 MB limit")
    try:
        data = parse_csv(content.decode("utf-8-sig"), units)
    except UnicodeDecodeError:
        raise HTTPException(422, "CSV must use UTF-8 encoding")
    with Session() as db:
        s, _, _ = get_scenario(db, request, sid)
        if not data["errors"]:
            try:
                Inputs.model_validate({**s.inputs, "blocks": data["blocks"]})
            except ValidationError as exc:
                data["errors"].append({"row": 0, "message": str(exc)[:800]})
    return data


class CommitImport(BaseModel):
    revision: int = Field(ge=1)
    content: str = Field(max_length=2_000_000)
    units: str = "m"


@app.post("/api/scenarios/{sid}/import")
def commit_import(sid: str, body: CommitImport, request: Request):
    data = parse_csv(body.content, body.units)
    if data["errors"]:
        raise HTTPException(
            422, {"message": "Import contains rejected rows", "errors": data["errors"]}
        )
    with Session.begin() as db:
        s, u, _ = get_scenario(db, request, sid, True)
        inputs = {
            **s.inputs,
            "blocks": data["blocks"],
            "source": {
                "kind": "csv",
                "hash": data["source_hash"],
                "coordinate_input_unit": body.units,
                "density_unit": "t/m3",
                "original_csv": body.content,
            },
        }
        return replace_inputs(db, s, u, body.revision, inputs, "model.imported")


class SurfaceImport(BaseModel):
    revision: int = Field(ge=1)
    content: str = Field(max_length=2_000_000)
    name: str = Field(min_length=1, max_length=120)
    units: str = "m"


@app.post("/api/scenarios/{sid}/surface-preview")
def surface_preview(sid: str, body: SurfaceImport, request: Request):
    with Session() as db:
        get_scenario(db, request, sid, True)
    mesh = parse_obj(body.content, body.units, body.name)
    return {
        "surface": mesh,
        "errors": [],
        "note": "Topology validated for viewing only. No solid volume is derived, even for closed meshes.",
    }


@app.post("/api/scenarios/{sid}/surface")
def surface_commit(sid: str, body: SurfaceImport, request: Request):
    mesh = parse_obj(body.content, body.units, body.name)
    with Session.begin() as db:
        s, u, _ = get_scenario(db, request, sid, True)
        old = s.inputs.get("surfaces", [])
        if any(m["id"] == mesh["id"] for m in old):
            raise HTTPException(409, "This surface is already imported")
        inputs = {**s.inputs, "surfaces": [*old, mesh]}
        return replace_inputs(db, s, u, body.revision, inputs, "surface.imported")


@app.get("/api/scenarios/{sid}/precheck")
def precheck(sid: str, request: Request):
    with Session() as db:
        s, _, _ = get_scenario(db, request, sid)
        return {
            "issues": graph_issues(Inputs.model_validate(s.inputs)),
            "activities": len(s.inputs["blocks"]),
        }


class NewRun(BaseModel):
    revision: int = Field(ge=1)
    idempotency_key: str = Field(min_length=8, max_length=100)


@app.post("/api/scenarios/{sid}/runs", status_code=202)
def submit(sid: str, body: NewRun, request: Request):
    with Session.begin() as db:
        # Serialize project submissions: idempotency and active-job checks are atomic.
        from .db import engine

        if engine.dialect.name == "sqlite":
            db.connection().exec_driver_sql("BEGIN IMMEDIATE")
        s, u, _ = get_scenario(db, request, sid, True)
        if engine.dialect.name == "postgresql":
            db.execute(select(Project.id).where(Project.id == s.project_id).with_for_update())
            db.refresh(s)
        previous = db.scalar(
            select(Run).where(Run.scenario_id == sid, Run.idempotency_key == body.idempotency_key)
        )
        if previous:
            return run_data(previous, False)
        if s.revision != body.revision:
            raise HTTPException(409, "Inputs changed; reload before scheduling")
        cfg = Inputs.model_validate(s.inputs)
        if not cfg.blocks:
            raise HTTPException(422, "Import at least one block before scheduling")
        issues = graph_issues(cfg)
        if issues:
            raise HTTPException(422, "; ".join(issues))
        if db.scalar(
            select(Run.id).where(
                Run.project_id == s.project_id, Run.status.in_(["queued", "running"])
            )
        ):
            raise HTTPException(409, "A schedule job is already active for this project")
        snapshot = copy.deepcopy(s.inputs)
        project = db.get(Project, s.project_id)
        snapshot["source"] = {
            **snapshot["source"],
            "project_reference": {
                "crs": project.crs,
                "vertical_datum": project.datum,
                "coordinate_unit": "m",
            },
        }
        r = Run(
            project_id=s.project_id,
            scenario_id=s.id,
            revision=s.revision,
            created_by=u.id,
            idempotency_key=body.idempotency_key,
            input_hash=input_hash(snapshot),
            snapshot=snapshot,
        )
        db.add(r)
        db.flush()
        audit(db, s.project_id, u.name, "run.queued", r.id)
        return run_data(r, False)


@app.get("/api/projects/{pid}/runs")
def runs(pid: str, request: Request):
    with Session() as db:
        access(db, request, pid)
        return [
            run_data(r, False)
            for r in db.scalars(
                select(Run).where(Run.project_id == pid).order_by(Run.created.desc()).limit(100)
            )
        ]


@app.get("/api/runs/{rid}")
def run(rid: str, request: Request):
    with Session() as db:
        r = db.get(Run, rid)
        if not r:
            raise HTTPException(404, "Run not found")
        access(db, request, r.project_id)
        return run_data(r)


@app.post("/api/runs/{rid}/cancel")
def cancel(rid: str, request: Request):
    with Session.begin() as db:
        r = db.get(Run, rid)
        if not r:
            raise HTTPException(404, "Run not found")
        u, _ = access(db, request, r.project_id, edit=True)
        if r.status not in ["queued", "running"]:
            raise HTTPException(409, "Run is already finished")
        r.cancel_requested = 1
        if r.status == "queued":
            r.status = "cancelled"
        audit(db, r.project_id, u.name, "run.cancelled", rid)
        return {"ok": True}


@app.post("/api/runs/{rid}/approve")
def approve(rid: str, request: Request):
    with Session.begin() as db:
        r = db.get(Run, rid)
        if not r:
            raise HTTPException(404, "Run not found")
        u, _ = access(db, request, r.project_id, approve=True)
        if r.status != "feasible" or not r.result or validate_result(r.snapshot, r.result):
            raise HTTPException(409, "Only independently validated feasible runs can be approved")
        if not r.approved_at:
            r.approved_at = now()
            r.approved_by = u.name
            audit(db, r.project_id, u.name, "run.approved", rid)
        return run_data(r, False)


class Membership(BaseModel):
    email: str = Field(max_length=200)
    role: str


@app.get("/api/projects/{pid}/members")
def members(pid: str, request: Request):
    with Session() as db:
        access(db, request, pid)
        return [
            {"name": u.name, "email": u.email, "role": m.role}
            for m, u in db.execute(
                select(Member, User)
                .join(User, User.id == Member.user_id)
                .where(Member.project_id == pid)
            )
        ]


@app.post("/api/projects/{pid}/members")
def add_member(pid: str, body: Membership, request: Request):
    if body.role not in ["planner", "viewer", "reviewer"]:
        raise HTTPException(422, "Choose planner, reviewer or viewer")
    with Session.begin() as db:
        u, _ = access(db, request, pid, owner=True)
        target = db.scalar(select(User).where(User.email == body.email.strip().lower()))
        if not target:
            raise HTTPException(
                422,
                "Account not found. Create an account using the server user-management command first.",
            )
        m = db.scalar(select(Member).where(Member.project_id == pid, Member.user_id == target.id))
        if m and m.role == "owner":
            raise HTTPException(409, "Project owner role cannot be replaced")
        if m:
            m.role = body.role
        else:
            db.add(Member(project_id=pid, user_id=target.id, role=body.role))
        audit(db, pid, u.name, "membership.changed", f"{target.email}: {body.role}")
        return {"ok": True}


@app.get("/api/projects/{pid}/audit")
def audits(pid: str, request: Request):
    with Session() as db:
        access(db, request, pid)
        return [
            {
                "id": a.id,
                "actor": a.actor,
                "action": a.action,
                "detail": a.detail,
                "created": a.created,
            }
            for a in db.scalars(
                select(Audit)
                .where(Audit.project_id == pid)
                .order_by(Audit.created.desc())
                .limit(150)
            )
        ]


EXPORT_FIELDS = [
    "block_id",
    "resource_id",
    "destination",
    "start",
    "end",
    "tonnes",
    "hours",
    "material",
    "basis",
    "ash",
]


def safe_cell(value):
    s = str(value) if value is not None else ""
    return "'" + s if s.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")) else s


@app.get("/api/runs/{rid}/export/{format}")
def export(rid: str, format: str, request: Request):
    with Session.begin() as db:
        r = db.get(Run, rid)
        if not r:
            raise HTTPException(404, "Run not found")
        u, _ = access(db, request, r.project_id)
        if not r.result:
            raise HTTPException(409, "No completed result to export")
        audit(db, r.project_id, u.name, "run.exported", f"{rid} {format}")
        manifest = {
            "run_id": r.id,
            "input_hash": r.input_hash,
            "scenario_revision": r.revision,
            "status": r.status,
            "engine": r.result["engine"],
            "created": r.created,
            "approved_by": r.approved_by,
            "timezone": r.snapshot["timezone"],
            "reference_system": r.snapshot.get("source", {}).get("project_reference", {}),
            "units": "tonnes; working hours; in-situ m3; dry/as_received as labelled",
            "assumptions": "Splittable mining tasks; same-day shifts; no stock reclaim; heuristic baseline, not global optimisation",
        }
        rows = r.result["allocations"]
        if format == "json":
            payload = json.dumps(
                {"manifest": manifest, "snapshot": r.snapshot, "result": r.result}, indent=2
            ).encode()
            mime = "application/json"
            ext = "json"
        elif format == "csv":
            out = io.StringIO()
            w = csv.writer(out)
            w.writerow(EXPORT_FIELDS + ["run_id", "input_hash", "scenario_revision"])
            for a in rows:
                w.writerow(
                    [safe_cell(a.get(k)) for k in EXPORT_FIELDS] + [rid, r.input_hash, r.revision]
                )
            payload = out.getvalue().encode("utf-8-sig")
            mime = "text/csv"
            ext = "csv"
        elif format == "xlsx":
            from openpyxl import Workbook

            wb = Workbook()
            ws = wb.active
            ws.title = "Allocations"
            ws.append(EXPORT_FIELDS)
            for a in rows:
                ws.append(
                    [
                        safe_cell(a.get(k)) if isinstance(a.get(k), str) else a.get(k)
                        for k in EXPORT_FIELDS
                    ]
                )
            meta = wb.create_sheet("Provenance")
            meta.append(["Field", "Value"])
            for k, v in manifest.items():
                meta.append([k, safe_cell(v)])
            totals = wb.create_sheet("Totals")
            totals.append(["Quantity", "Value"])
            for k, v in r.result["totals"].items():
                totals.append([k, v])
            out = io.BytesIO()
            wb.save(out)
            payload = out.getvalue()
            mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ext = "xlsx"
        elif format == "pdf":
            from reportlab.pdfgen import canvas
            from reportlab.lib.pagesizes import A4

            out = io.BytesIO()
            c = canvas.Canvas(out, pagesize=A4)
            y = 800

            def line(s):
                nonlocal y
                if y < 55:
                    c.showPage()
                    c.setFont("Helvetica", 9)
                    y = 800
                from textwrap import wrap

                chunks = wrap(str(s), 100, replace_whitespace=False) or [""]
                for chunk in chunks:
                    if y < 55:
                        c.showPage()
                        c.setFont("Helvetica", 9)
                        y = 800
                    c.drawString(40, y, chunk)
                    y -= 17

            c.setFont("Helvetica-Bold", 18)
            line("MESSI | Production plan")
            c.setFont("Helvetica", 9)
            for k, v in manifest.items():
                line(f"{k}: {v}")
            line("")
            line("RECONCILED PHYSICAL TOTALS")
            for k, v in r.result["totals"].items():
                line(f"{k}: {v:,.3f}")
            line("")
            line("ACTIVITY / START UTC / END UTC / TONNES")
            for a in r.result["activities"]:
                line(
                    f"{a['id']} | {a['start'] or 'unscheduled'} | {a['end'] or '-'} | {a['tonnes']:,.3f}"
                )
            for u in r.result["unscheduled"]:
                line(f"Unscheduled {u['id']}: {u['reason']}")
            c.save()
            payload = out.getvalue()
            mime = "application/pdf"
            ext = "pdf"
        else:
            raise HTTPException(404, "Choose csv, xlsx, pdf or json")
        return StreamingResponse(
            io.BytesIO(payload),
            media_type=mime,
            headers={"Content-Disposition": f'attachment; filename="MESSI-{rid[:8]}.{ext}"'},
        )


DIST = ROOT / "frontend" / "dist"
if DIST.exists():
    app.mount("/assets", StaticFiles(directory=DIST / "assets"), name="assets")

    @app.get("/{path:path}")
    def spa(path: str):
        if path.startswith("api/"):
            raise HTTPException(404, "API route not found")
        return FileResponse(DIST / "index.html")
